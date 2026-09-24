import tkinter as tk
from tkinter import ttk, scrolledtext, filedialog, messagebox
import pandas as pd
import numpy as np
import random
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

# pip install pandas openpyxl numpy


SCORING_ITEMS = {
    "礼貌": 5,
    "床上": 10,
    "床下": 10,
    "个人衣服": 10,
    "地面卫生": 10,
    "电线": 10,
    "窗台": 5,
    "乱挂杂物字画": 10,
    "桌椅": 5,
    "桌面": 10,
    "厕所": 10,
    "及时清理垃圾": 5,
}

REPORT_ROWS = [
    ("整体", "礼貌情况(5分)", "礼貌"),
    ("个人", "床上物品被子、被单、\n枕头摆放整齐(10分)", "床上"),
    (None, "床下鞋子、行李及物品\n摆放整齐(10分)", "床下"),
    (None, "个人衣物及洗漱用品\n摆放整齐(10分)", "个人衣服"),
    ("地面", "干净无脏水纸屑等杂\n物(10分)", "地面卫生"),
    ("墙壁", "严禁电网线乱挂或纠\n缠(10分)", "电线"),
    (None, "窗台无垃圾杂物\n(5分)", "窗台"),
    (None, "严禁乱挂杂物，乱贴字\n画(10分)", "乱挂杂物字画"),
    ("桌椅", "桌椅摆放整齐(5分)", "桌椅"),
    (None, "桌面物品摆放整齐(10\n分)", "桌面"),
    ("卫生间", "卫生间干净无异味(10\n分)", "厕所"),
    (None, "及时清理垃圾(5分)", "及时清理垃圾"),
]


def export_inspection_workbook(results, output_path, inspection_date, grade, inspector):
    """Write generated scores in the horizontal college inspection form layout."""
    if not results:
        raise ValueError("没有可导出的寝室评分")

    dorm_count = len(results)
    first_dorm_col = 3  # C
    last_col = first_dorm_col + dorm_count - 1
    last_col_letter = get_column_letter(last_col)
    compact_header = dorm_count < 7
    title_row = 2
    if compact_header:
        dorm_header_row = 6
        data_start_row = 7
    else:
        dorm_header_row = 4
        data_start_row = 5
    remark_row = data_start_row + len(REPORT_ROWS)
    total_row = remark_row + 1

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "院公检记录表"
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 30
    for col in range(first_dorm_col, last_col + 1):
        worksheet.column_dimensions[get_column_letter(col)].width = 12

    # The reference workbook uses Calibri; spreadsheet apps supply their local CJK fallback.
    title_font = Font(name="Calibri", size=14, bold=True)
    body_font = Font(name="Calibri", size=10)
    table_font = Font(name="Calibri", size=9)
    thin_side = Side(style="thin", color="000000")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)

    worksheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=last_col)
    title = worksheet.cell(title_row, 1, "管理学院寝室卫生检查")
    title.font = title_font
    title.alignment = centered
    worksheet.row_dimensions[title_row].height = 28

    if compact_header:
        metadata = (
            f"检查日期: {inspection_date}",
            f"所属年级: {grade}",
            f"检查人员: {inspector}",
        )
        for row, value in enumerate(metadata, start=3):
            worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
            cell = worksheet.cell(row, 1, value)
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            worksheet.row_dimensions[row].height = 21
    else:
        # Reserve the first two columns for the dormitory label and row group.
        worksheet.merge_cells("A3:B3")
        worksheet["A3"] = "宿舍号\n要求 评分"
        worksheet["A3"].font = body_font
        worksheet["A3"].alignment = centered
        metadata_ranges = (
            (3, 3, 4, f"日期:{inspection_date}"),
            (3, 5, 8, f"所属年级:{grade}"),
            (3, 9, last_col, f"检查人员:{inspector}"),
        )
        for row, start_col, end_col, value in metadata_ranges:
            if start_col > last_col:
                continue
            end_col = min(end_col, last_col)
            if end_col > start_col:
                worksheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
            worksheet.cell(row, start_col, value)
        for cell in worksheet[3]:
            if 3 <= cell.column <= last_col and cell.value is not None:
                cell.font = body_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        worksheet.row_dimensions[3].height = 24

    # Column headings and scoring criteria.
    if compact_header:
        worksheet.merge_cells(start_row=dorm_header_row, start_column=1, end_row=dorm_header_row,
                              end_column=2)
        worksheet.cell(dorm_header_row, 1, "宿舍号\n要求 评分")

    for col in range(1, last_col + 1):
        cell = worksheet.cell(dorm_header_row, col)
        cell.font = table_font
        cell.alignment = centered
        cell.border = cell_border
    for index, result in enumerate(results):
        dorm_cell = worksheet.cell(dorm_header_row, first_dorm_col + index, str(result["寝室"]))
        dorm_cell.number_format = "@"

    worksheet.row_dimensions[dorm_header_row].height = 34
    for index, (group, description, score_key) in enumerate(REPORT_ROWS):
        row = data_start_row + index
        if group is not None:
            worksheet.cell(row, 1, group)
        worksheet.cell(row, 2, description)
        worksheet.row_dimensions[row].height = 34 if "\n" in description else 24
        for col in range(1, last_col + 1):
            cell = worksheet.cell(row, col)
            cell.font = table_font
            cell.alignment = centered
            cell.border = cell_border
        for index, result in enumerate(results):
            cell = worksheet.cell(row, first_dorm_col + index, result[score_key])
            cell.number_format = "0"

    # Group labels follow the merged vertical sections in the reference form.
    for start_row, end_row in ((data_start_row + 1, data_start_row + 3),
                               (data_start_row + 5, data_start_row + 7),
                               (data_start_row + 8, data_start_row + 9),
                               (data_start_row + 10, data_start_row + 11)):
        worksheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
        cell = worksheet.cell(start_row, 1)
        cell.alignment = centered
        cell.border = cell_border

    worksheet.merge_cells(start_row=remark_row, start_column=1, end_row=remark_row, end_column=2)
    worksheet.cell(remark_row, 1, "备注(有无拒检现象)")
    worksheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    worksheet.cell(total_row, 1, "总分(100分)")
    for row in (remark_row, total_row):
        worksheet.row_dimensions[row].height = 24
        for col in range(1, last_col + 1):
            cell = worksheet.cell(row, col)
            cell.font = table_font
            cell.alignment = centered
            cell.border = cell_border
    for index, result in enumerate(results):
        col = first_dorm_col + index
        worksheet.cell(remark_row, col, "无")
        total_cell = worksheet.cell(total_row, col, result["总分"])
        total_cell.number_format = "0"

    worksheet.freeze_panes = worksheet.cell(dorm_header_row + 1, first_dorm_col)
    worksheet.print_area = f"A{title_row}:{last_col_letter}{total_row}"
    worksheet.page_setup.orientation = "landscape" if dorm_count > 11 else "portrait"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(output_path)

class UltimateDormScoreApp:
    def __init__(self, root):
        self.root = root
        self.root.title("寝室卫生检查评分系统")
        self.root.geometry("780x900")
        self.root.minsize(700, 780)
        
        self.setup_styles()
        self.main_frame = ttk.Frame(root, padding=(28, 24), style="App.TFrame")
        self.main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        root.configure(background="#F4F6F7")
        root.columnconfigure(0, weight=1)
        root.rowconfigure(0, weight=1)
        self.main_frame.columnconfigure(0, weight=1)
        self.main_frame.rowconfigure(3, weight=1)
        
        self.create_widgets()
        self.restore_default_dorms()

    def setup_styles(self):
        style = ttk.Style()
        style.theme_use("clam")

        background = "#F4F6F7"
        surface = "#FFFFFF"
        border = "#E1E5E8"
        text = "#20262D"
        muted = "#75808A"
        accent = "#34414D"

        style.configure("App.TFrame", background=background)
        style.configure("Surface.TFrame", background=surface)
        style.configure(
            "Card.TFrame",
            background=surface,
            bordercolor=border,
            lightcolor=border,
            darkcolor=border,
            borderwidth=1,
            relief="solid",
        )
        style.configure("TLabel", background=background, foreground=text, font=("微软雅黑", 10))
        style.configure("Card.TLabel", background=surface, foreground=text, font=("微软雅黑", 10))
        style.configure(
            "Title.TLabel",
            background=background,
            foreground=text,
            font=("微软雅黑", 22, "bold"),
        )
        style.configure("Subtitle.TLabel", background=background, foreground=muted, font=("微软雅黑", 10))
        style.configure("Section.TLabel", background=surface, foreground=text, font=("微软雅黑", 11, "bold"))
        style.configure("Field.TLabel", background=surface, foreground=muted, font=("微软雅黑", 9))
        style.configure("Hint.TLabel", background=surface, foreground=muted, font=("微软雅黑", 9))
        style.configure("Status.TLabel", background=background, foreground=muted, font=("微软雅黑", 9))
        style.configure(
            "TEntry",
            fieldbackground=surface,
            foreground=text,
            padding=(9, 7),
            bordercolor=border,
            lightcolor=border,
            darkcolor=border,
        )
        style.map("TEntry", bordercolor=[("focus", accent)], lightcolor=[("focus", accent)])
        style.configure(
            "Quiet.TButton",
            background=surface,
            foreground=text,
            padding=(12, 8),
            bordercolor=border,
            font=("微软雅黑", 9),
        )
        style.map(
            "Quiet.TButton",
            background=[("active", "#EEF1F3"), ("pressed", "#E7EBEE")],
            bordercolor=[("focus", border)],
        )
        style.configure(
            "Primary.TButton",
            background=accent,
            foreground=surface,
            padding=(20, 9),
            borderwidth=0,
            font=("微软雅黑", 10, "bold"),
        )
        style.map(
            "Primary.TButton",
            background=[("active", "#465563"), ("pressed", "#29343E")],
            foreground=[("disabled", "#D9DEE2")],
        )

    def get_default_dorms(self):
        return "\n".join(["10A201", "10B408", "10B414"])

    def get_current_date(self):
        return datetime.now().strftime("%y.%m.%d")

    def create_widgets(self):
        header = ttk.Frame(self.main_frame, style="App.TFrame")
        header.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 18))
        ttk.Label(header, text="查寝小帮手", style="Title.TLabel").pack(anchor=tk.W)
        ttk.Label(header, text="管理学院 · 院公检记录表", style="Subtitle.TLabel").pack(anchor=tk.W, pady=(3, 0))

        info_card = ttk.Frame(self.main_frame, style="Card.TFrame", padding=(18, 15))
        info_card.grid(row=1, column=0, sticky=(tk.W, tk.E), pady=(0, 12))
        info_card.columnconfigure(0, weight=1, uniform="info")
        info_card.columnconfigure(1, weight=1, uniform="info")
        ttk.Label(info_card, text="检查信息", style="Section.TLabel").grid(
            row=0, column=0, columnspan=2, sticky=tk.W, pady=(0, 12)
        )

        self.date_var = tk.StringVar(value=self.get_current_date())
        self.inspector_var = tk.StringVar()
        self.grade_var = tk.StringVar()
        self.seed_var = tk.StringVar(value="666")
        fields = (
            ("检查日期", self.date_var),
            ("检查人员", self.inspector_var),
            ("所属年级", self.grade_var),
            ("随机种子（相同种子结果一致）", self.seed_var),
        )
        for index, (label, variable) in enumerate(fields):
            field = ttk.Frame(info_card, style="Surface.TFrame")
            field.grid(row=1 + index // 2, column=index % 2, sticky=(tk.W, tk.E),
                       padx=(0, 12) if index % 2 == 0 else (12, 0), pady=(0, 10))
            field.columnconfigure(0, weight=1)
            ttk.Label(field, text=label, style="Field.TLabel").grid(row=0, column=0, sticky=tk.W, pady=(0, 5))
            ttk.Entry(field, textvariable=variable).grid(row=1, column=0, sticky=(tk.W, tk.E))

        score_card = ttk.Frame(self.main_frame, style="Card.TFrame", padding=(18, 14))
        score_card.grid(row=2, column=0, sticky=(tk.W, tk.E), pady=(0, 12))
        score_card.columnconfigure(0, weight=1)
        ttk.Label(score_card, text="评分范围", style="Section.TLabel").grid(row=0, column=0, sticky=tk.W)
        ttk.Label(score_card, text="按正态分布生成目标总分", style="Hint.TLabel").grid(
            row=0, column=1, sticky=tk.E
        )

        self.min_score_var = tk.StringVar(value="84")
        self.max_score_var = tk.StringVar(value="96")
        range_fields = ttk.Frame(score_card, style="Surface.TFrame")
        range_fields.grid(row=1, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=(10, 0))
        range_fields.columnconfigure(0, weight=1, uniform="range")
        range_fields.columnconfigure(2, weight=1, uniform="range")
        ttk.Label(range_fields, text="最低分", style="Field.TLabel").grid(row=0, column=0, sticky=tk.W, pady=(0, 5))
        ttk.Label(range_fields, text="最高分", style="Field.TLabel").grid(row=0, column=2, sticky=tk.W, pady=(0, 5))
        ttk.Entry(range_fields, textvariable=self.min_score_var, width=10).grid(row=1, column=0, sticky=tk.W)
        ttk.Label(range_fields, text="—", style="Hint.TLabel").grid(row=1, column=1, padx=18)
        ttk.Entry(range_fields, textvariable=self.max_score_var, width=10).grid(row=1, column=2, sticky=tk.W)

        dorm_card = ttk.Frame(self.main_frame, style="Card.TFrame", padding=(18, 14))
        dorm_card.grid(row=3, column=0, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(0, 12))
        dorm_card.columnconfigure(0, weight=1)
        dorm_card.rowconfigure(1, weight=1)
        ttk.Label(dorm_card, text="寝室名单", style="Section.TLabel").grid(row=0, column=0, sticky=tk.W)
        ttk.Label(dorm_card, text="每行填写一个寝室号", style="Hint.TLabel").grid(
            row=0, column=1, sticky=tk.E
        )
        self.dorm_text = scrolledtext.ScrolledText(
            dorm_card,
            width=50,
            height=10,
            wrap=tk.WORD,
            font=("TkFixedFont", 11),
            background="#FFFFFF",
            foreground="#20262D",
            insertbackground="#34414D",
            selectbackground="#DCE3E8",
            selectforeground="#20262D",
            relief=tk.FLAT,
            borderwidth=0,
            highlightthickness=1,
            highlightbackground="#E1E5E8",
            highlightcolor="#34414D",
            padx=11,
            pady=10,
        )
        self.dorm_text.grid(row=1, column=0, columnspan=2, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(10, 0))
        self.dorm_text.vbar.configure(
            background="#F4F6F7",
            troughcolor="#FFFFFF",
            activebackground="#D9DEE2",
            relief=tk.FLAT,
            borderwidth=0,
            highlightthickness=0,
        )

        path_card = ttk.Frame(self.main_frame, style="Card.TFrame", padding=(18, 13))
        path_card.grid(row=4, column=0, sticky=(tk.W, tk.E), pady=(0, 14))
        path_card.columnconfigure(0, weight=1)
        ttk.Label(path_card, text="导出位置", style="Section.TLabel").grid(row=0, column=0, columnspan=2,
                                                                          sticky=tk.W, pady=(0, 8))
        self.path_var = tk.StringVar(value=os.path.join(os.getcwd(), "寝室卫生检查评分表.xlsx"))
        ttk.Entry(path_card, textvariable=self.path_var).grid(row=1, column=0, sticky=(tk.W, tk.E), padx=(0, 9))
        ttk.Button(path_card, text="浏览…", style="Quiet.TButton", command=self.browse_file).grid(row=1, column=1)

        actions = ttk.Frame(self.main_frame, style="App.TFrame")
        actions.grid(row=5, column=0, sticky=(tk.W, tk.E))
        actions.columnconfigure(1, weight=1)
        ttk.Button(actions, text="恢复默认寝室", style="Quiet.TButton",
                   command=self.restore_default_dorms).grid(row=0, column=0, sticky=tk.W)
        self.status_var = tk.StringVar(value="准备就绪")
        ttk.Label(actions, textvariable=self.status_var, style="Status.TLabel", anchor=tk.E).grid(
            row=0, column=1, sticky=(tk.W, tk.E), padx=12
        )
        ttk.Button(actions, text="生成评分表格", style="Primary.TButton",
                   command=self.generate_scores).grid(row=0, column=2, sticky=tk.E)

    def restore_default_dorms(self):
        self.dorm_text.delete("1.0", tk.END)
        self.dorm_text.insert("1.0", self.get_default_dorms())

    def browse_file(self):
        file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if file_path: self.path_var.set(file_path)

    def generate_scores(self):
        try:
            # 1. 数据校验
            min_s = int(self.min_score_var.get())
            max_s = int(self.max_score_var.get())
            if min_s >= max_s: raise ValueError("最低分必须小于最高分")

            # 2. 初始化环境
            seed_val = int(self.seed_var.get()) if self.seed_var.get() else None
            np.random.seed(seed_val)
            random.seed(seed_val)

            scoring_items = SCORING_ITEMS.copy()
            
            dorms = [line.strip() for line in self.dorm_text.get("1.0", tk.END).split('\n') if line.strip()]
            if not dorms: raise ValueError("寝室列表不能为空")

            # 3. 正态分布参数
            mu = (min_s + max_s) / 2
            sigma = (max_s - min_s) / 6
            
            results = []
            for dorm in dorms:
                # 目标总分生成
                target = np.clip(np.random.normal(mu, sigma), min_s, max_s)
                
                # 初始化分值 (礼貌固定满分)
                scores = {"日期": self.date_var.get(), "检查人员": self.inspector_var.get(), "寝室": dorm, "礼貌": 5}
                current_sum = 5
                
                # 其他项随机赋初值
                other_keys = [k for k in scoring_items.keys() if k != "礼貌"]
                for k in other_keys:
                    val = random.choice([3, 5]) if scoring_items[k] == 5 else random.choice([6, 8, 10])
                    scores[k] = val
                    current_sum += val
                
                # 算法逼近
                for _ in range(50):
                    if abs(current_sum - target) <= 1: break
                    item = random.choice(other_keys)
                    step = 2
                    if current_sum < target and scores[item] + step <= scoring_items[item]:
                        scores[item] += step
                        current_sum += step
                    elif current_sum > target and scores[item] - step >= (6 if scoring_items[item]==10 else 3):
                        scores[item] -= step
                        current_sum -= step
                
                scores["总分"] = current_sum
                results.append(scores)

            # 4. 保存与反馈
            df = pd.DataFrame(results)
            export_inspection_workbook(
                results,
                self.path_var.get(),
                self.date_var.get(),
                self.grade_var.get(),
                self.inspector_var.get(),
            )
            
            stats = f"成功生成！\n平均分: {df['总分'].mean():.2f}\n标准差: {df['总分'].std():.2f}\n\n各项均分：\n"
            for k in scoring_items.keys():
                stats += f"{k}: {df[k].mean():.2f}  "
            
            messagebox.showinfo("完成", stats)
            self.status_var.set(f"已保存至: {self.path_var.get()}")

        except Exception as e:
            messagebox.showerror("错误", str(e))

if __name__ == "__main__":
    root = tk.Tk()
    app = UltimateDormScoreApp(root)
    root.mainloop()
